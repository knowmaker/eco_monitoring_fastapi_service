from dataclasses import dataclass
from datetime import datetime
from io import BytesIO

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.core.dates import APP_TIMEZONE, bucket_label, parse_local_datetime, to_epoch_ms
from app.core.values import to_float
from app.models.cagg_dust_daily import CaggDustDaily
from app.models.cagg_dust_hourly import CaggDustHourly
from app.models.cagg_gas_daily import CaggGasDaily
from app.models.cagg_gas_hourly import CaggGasHourly
from app.models.cagg_ivtm_daily import CaggIvtmDaily
from app.models.cagg_ivtm_hourly import CaggIvtmHourly
from app.models.cagg_meteo_daily import CaggMeteoDaily
from app.models.cagg_meteo_hourly import CaggMeteoHourly
from app.models.cagg_profile_inversion_daily import CaggProfileInversionDaily
from app.models.cagg_profile_inversion_hourly import CaggProfileInversionHourly
from app.models.cagg_profile_levels_daily import CaggProfileLevelsDaily
from app.models.cagg_profile_levels_hourly import CaggProfileLevelsHourly
from app.models.monitoring_posts import MonitoringPost
from app.schemas.export_aggregates import ExportAggregatesRequest


MAX_EXPORT_ROWS = 250_000
GAS_SUBSTANCE_ORDER = ["NO2", "O3", "NO", "SO2", "CO", "H2S"]


@dataclass(frozen=True)
class ExportFile:
    content: bytes
    filename: str


def build_aggregates_export(db: Session, payload: ExportAggregatesRequest) -> ExportFile:
    start_local, end_local = _parse_export_period(payload)
    _validate_device_types(payload.device_types)

    station_rows = _resolve_stations(db, payload.station_ids)
    if not station_rows:
        raise HTTPException(status_code=422, detail="No confirmed stations available for export")

    start_ms = to_epoch_ms(start_local)
    end_ms = to_epoch_ms(end_local)
    station_ids = [row.id for row in station_rows]
    models = _models_by_device_type(payload.aggregation)

    workbook = Workbook()
    _append_info_sheet(workbook, payload, station_rows, start_local, end_local)
    row_count = 0

    if "gas" in payload.device_types:
        row_count = _append_gas_sheet(
            workbook, db, models["gas"], station_ids, start_ms, end_ms, payload.aggregation, row_count
        )
    if "dust" in payload.device_types:
        row_count = _append_dust_sheet(
            workbook, db, models["dust"], station_ids, start_ms, end_ms, payload.aggregation, row_count
        )
    if "meteo" in payload.device_types:
        row_count = _append_meteo_sheet(
            workbook, db, models["meteo"], station_ids, start_ms, end_ms, payload.aggregation, row_count
        )
    if "ivtm" in payload.device_types:
        row_count = _append_ivtm_sheet(
            workbook, db, models["ivtm"], station_ids, start_ms, end_ms, payload.aggregation, row_count
        )
    if "profile" in payload.device_types:
        row_count = _append_profile_levels_sheet(
            workbook,
            db,
            models["profile_levels"],
            station_ids,
            start_ms,
            end_ms,
            payload.aggregation,
            row_count,
        )
        row_count = _append_profile_inversion_sheet(
            workbook,
            db,
            models["profile_inversion"],
            station_ids,
            start_ms,
            end_ms,
            payload.aggregation,
            row_count,
        )

    output = BytesIO()
    workbook.save(output)
    return ExportFile(
        content=output.getvalue(),
        filename=_build_filename(payload, start_local, end_local),
    )


def _parse_export_period(payload: ExportAggregatesRequest) -> tuple[datetime, datetime]:
    try:
        start_local = parse_local_datetime(payload.start, is_end=False)
        end_local = parse_local_datetime(payload.end, is_end=True)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(exc)) from exc

    if end_local < start_local:
        raise HTTPException(status_code=422, detail="end must be greater than or equal to start")
    return start_local, end_local


def _validate_device_types(device_types: list[str]) -> None:
    if not device_types:
        raise HTTPException(status_code=422, detail="Select at least one device type")
    if len(set(device_types)) != len(device_types):
        raise HTTPException(status_code=422, detail="Device types must be unique")


def _resolve_stations(db: Session, station_ids: list[int] | None) -> list[MonitoringPost]:
    if station_ids is None:
        return db.scalars(
            select(MonitoringPost)
            .where(MonitoringPost.is_confirmed.is_(True))
            .order_by(MonitoringPost.serial.asc())
        ).all()

    if not station_ids:
        raise HTTPException(status_code=422, detail="Select at least one station")

    station_rows = db.scalars(
        select(MonitoringPost)
        .where(
            MonitoringPost.id.in_(station_ids),
            MonitoringPost.is_confirmed.is_(True),
        )
        .order_by(MonitoringPost.serial.asc())
    ).all()
    found_ids = {row.id for row in station_rows}
    if set(station_ids) - found_ids:
        raise HTTPException(status_code=404, detail="One or more stations were not found")
    return station_rows


def _models_by_device_type(aggregation: str) -> dict[str, type]:
    if aggregation == "hourly":
        return {
            "gas": CaggGasHourly,
            "dust": CaggDustHourly,
            "meteo": CaggMeteoHourly,
            "ivtm": CaggIvtmHourly,
            "profile_levels": CaggProfileLevelsHourly,
            "profile_inversion": CaggProfileInversionHourly,
        }
    return {
        "gas": CaggGasDaily,
        "dust": CaggDustDaily,
        "meteo": CaggMeteoDaily,
        "ivtm": CaggIvtmDaily,
        "profile_levels": CaggProfileLevelsDaily,
        "profile_inversion": CaggProfileInversionDaily,
    }


def _append_header(sheet: Worksheet, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def _append_info_sheet(
    workbook: Workbook,
    payload: ExportAggregatesRequest,
    station_rows: list[MonitoringPost],
    start_local: datetime,
    end_local: datetime,
) -> None:
    sheet = workbook.active
    sheet.title = "info"
    rows = [
        ("aggregation", payload.aggregation),
        ("timezone", APP_TIMEZONE),
        ("start", start_local.strftime("%Y-%m-%d %H:%M:%S")),
        ("end", end_local.strftime("%Y-%m-%d %H:%M:%S")),
        ("stations", ", ".join(post.name or "Без названия" for post in station_rows)),
        ("device_types", ", ".join(payload.device_types)),
    ]
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 60


def _base_query(model: type, station_ids: list[int], start_ms: int, end_ms: int):
    return (
        select(model, MonitoringPost.serial.label("post_serial"), MonitoringPost.name.label("post_name"))
        .join(MonitoringPost, MonitoringPost.id == model.monitoring_post_id)
        .where(
            model.monitoring_post_id.in_(station_ids),
            model.bucket_ms >= start_ms,
            model.bucket_ms <= end_ms,
        )
        .order_by(MonitoringPost.serial.asc(), model.bucket_ms.asc())
    )


def _sort_gas_codes(codes: list[str]) -> list[str]:
    known_codes = [code for code in GAS_SUBSTANCE_ORDER if code in codes]
    other_codes = sorted(code for code in codes if code not in GAS_SUBSTANCE_ORDER)
    return known_codes + other_codes


def _append_gas_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("gas")
    substance_rows = db.execute(
        select(model.substance_code)
        .where(
            model.monitoring_post_id.in_(station_ids),
            model.bucket_ms >= start_ms,
            model.bucket_ms <= end_ms,
            model.value_avg.is_not(None),
        )
        .distinct()
    )
    substance_codes = _sort_gas_codes(list({*GAS_SUBSTANCE_ORDER, *(row.substance_code for row in substance_rows)}))
    _append_header(sheet, ["station_serial", "station_name", "time", *substance_codes])

    rows = db.execute(
        _base_query(model, station_ids, start_ms, end_ms)
        .where(model.value_avg.is_not(None))
        .order_by(
            MonitoringPost.serial.asc(),
            model.bucket_ms.asc(),
            model.substance_code.asc(),
        )
    )
    current_key = None
    current_values: dict[str, float | None] = {}
    current_serial = None
    current_name = None
    current_bucket_ms = None

    for row in rows:
        item = row[0]
        key = (row.post_serial, item.bucket_ms)
        if current_key is not None and key != current_key:
            sheet.append([
                current_serial,
                current_name,
                bucket_label(current_bucket_ms, aggregation),
                *(current_values.get(code) for code in substance_codes),
            ])
            row_count = _increment_row_count(row_count)
            current_values = {}

        current_key = key
        current_serial = row.post_serial
        current_name = row.post_name
        current_bucket_ms = item.bucket_ms
        current_values[item.substance_code] = to_float(item.value_avg)

    if current_key is not None:
        sheet.append([
            current_serial,
            current_name,
            bucket_label(current_bucket_ms, aggregation),
            *(current_values.get(code) for code in substance_codes),
        ])
        row_count = _increment_row_count(row_count)

    return row_count


def _append_dust_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("dust")
    _append_header(sheet, [
        "station_serial",
        "station_name",
        "time",
        "pm1_mg_m3",
        "pm2_5_mg_m3",
        "pm10_mg_m3",
        "tsp_mg_m3",
    ])
    query = _base_query(model, station_ids, start_ms, end_ms).where(
        or_(
            model.pm1_avg.is_not(None),
            model.pm2_avg.is_not(None),
            model.pm10_avg.is_not(None),
            model.tsp_avg.is_not(None),
        )
    )
    for row in db.execute(query):
        item = row[0]
        sheet.append([
            row.post_serial,
            row.post_name,
            bucket_label(item.bucket_ms, aggregation),
            to_float(item.pm1_avg),
            to_float(item.pm2_avg),
            to_float(item.pm10_avg),
            to_float(item.tsp_avg),
        ])
        row_count = _increment_row_count(row_count)
    return row_count


def _append_meteo_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("meteo")
    _append_header(sheet, [
        "station_serial",
        "station_name",
        "time",
        "air_temperature_c",
        "air_humidity_percent",
        "atmospheric_pressure",
        "wind_direction_deg",
        "wind_speed_m_s",
    ])
    query = _base_query(model, station_ids, start_ms, end_ms).where(
        or_(
            model.air_temp_avg.is_not(None),
            model.air_hum_avg.is_not(None),
            model.atm_press_avg.is_not(None),
            model.hor_win_dir_avg.is_not(None),
            model.hor_win_spd_avg.is_not(None),
        )
    )
    for row in db.execute(query):
        item = row[0]
        sheet.append([
            row.post_serial,
            row.post_name,
            bucket_label(item.bucket_ms, aggregation),
            to_float(item.air_temp_avg),
            to_float(item.air_hum_avg),
            to_float(item.atm_press_avg),
            to_float(item.hor_win_dir_avg),
            to_float(item.hor_win_spd_avg),
        ])
        row_count = _increment_row_count(row_count)
    return row_count


def _append_ivtm_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("ivtm")
    _append_header(sheet, [
        "station_serial",
        "station_name",
        "time",
        "humidity_percent",
        "temperature_c",
    ])
    query = _base_query(model, station_ids, start_ms, end_ms).where(
        or_(
            model.sensor_ivtm_hum_avg.is_not(None),
            model.sensor_ivtm_temp_avg.is_not(None),
        )
    )
    for row in db.execute(query):
        item = row[0]
        sheet.append([
            row.post_serial,
            row.post_name,
            bucket_label(item.bucket_ms, aggregation),
            to_float(item.sensor_ivtm_hum_avg),
            to_float(item.sensor_ivtm_temp_avg),
        ])
        row_count = _increment_row_count(row_count)
    return row_count


def _append_profile_levels_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("profile_levels")
    _append_header(sheet, ["station_serial", "station_name", "time", "height_m", "temperature_c"])
    rows = db.execute(
        _base_query(model, station_ids, start_ms, end_ms)
        .where(model.temperature_avg.is_not(None))
        .order_by(
            MonitoringPost.serial.asc(),
            model.bucket_ms.asc(),
            model.height.asc(),
        )
    )
    for row in rows:
        item = row[0]
        sheet.append([
            row.post_serial,
            row.post_name,
            bucket_label(item.bucket_ms, aggregation),
            to_float(item.height),
            to_float(item.temperature_avg),
        ])
        row_count = _increment_row_count(row_count)
    return row_count


def _append_profile_inversion_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("profile_inversion")
    _append_header(sheet, [
        "station_serial",
        "station_name",
        "time",
        "inversion_power",
        "inversion_lower_m",
        "inversion_upper_m",
        "inversion_delta_t_c",
    ])
    query = _base_query(model, station_ids, start_ms, end_ms).where(
        or_(
            model.inversion_power_avg.is_not(None),
            model.inversion_lower_avg.is_not(None),
            model.inversion_upper_avg.is_not(None),
            model.inversion_delta_t_avg.is_not(None),
        )
    )
    for row in db.execute(query):
        item = row[0]
        sheet.append([
            row.post_serial,
            row.post_name,
            bucket_label(item.bucket_ms, aggregation),
            to_float(item.inversion_power_avg),
            to_float(item.inversion_lower_avg),
            to_float(item.inversion_upper_avg),
            to_float(item.inversion_delta_t_avg),
        ])
        row_count = _increment_row_count(row_count)
    return row_count


def _increment_row_count(row_count: int) -> int:
    row_count += 1
    if row_count > MAX_EXPORT_ROWS:
        raise HTTPException(status_code=413, detail="Слишком большой экспорт. Уменьшите диапазон или выборку.")
    return row_count


def _filename_datetime_part(value: datetime) -> str:
    return value.strftime("%Y-%m-%d_%H-%M")


def _build_filename(payload: ExportAggregatesRequest, start_local: datetime, end_local: datetime) -> str:
    period_part = f"{_filename_datetime_part(start_local)}_to_{_filename_datetime_part(end_local)}"
    return f"eco_export_{payload.aggregation}_{period_part}.xlsx"
