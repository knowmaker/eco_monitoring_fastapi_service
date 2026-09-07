from datetime import date, datetime, time
from io import BytesIO
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Response, status
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.db.session import get_db
from app.models.cagg_dust_daily import CaggDustDaily
from app.models.cagg_dust_hourly import CaggDustHourly
from app.models.cagg_gas_daily import CaggGasDaily
from app.models.cagg_gas_hourly import CaggGasHourly
from app.models.cagg_ivtm_daily import CaggIvtmDaily
from app.models.cagg_ivtm_hourly import CaggIvtmHourly
from app.models.cagg_meteo_daily import CaggMeteoDaily
from app.models.cagg_meteo_hourly import CaggMeteoHourly
from app.models.cagg_profile_inversion_daily import CaggProfileInversionDaily
from app.models.cagg_profile_inversion_hourly import CaggProfileInversionHourly
from app.models.cagg_profile_levels_daily import CaggProfileLevelsDaily
from app.models.cagg_profile_levels_hourly import CaggProfileLevelsHourly
from app.models.monitoring_posts import MonitoringPost
from app.models.user import User
from app.schemas.export_aggregates import ExportAggregatesRequest


APP_TIMEZONE = "Europe/Moscow"
MAX_EXPORT_ROWS = 250_000
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GAS_SUBSTANCE_ORDER = ["NO2", "O3", "NO", "SO2", "CO", "H2S"]

router = APIRouter(prefix="/export", tags=["export"])


def parse_local_datetime(value: str, *, is_end: bool) -> datetime:
    text = value.strip()
    try:
        if len(text) == 10:
            parsed_date = date.fromisoformat(text)
            parsed = datetime.combine(parsed_date, time.max if is_end else time.min)
        else:
            parsed = datetime.fromisoformat(text)
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="start and end must use ISO date or datetime format",
        ) from exc

    local_tz = ZoneInfo(APP_TIMEZONE)
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=local_tz)
    return parsed.astimezone(local_tz)


def to_epoch_ms(value: datetime) -> int:
    return int(value.timestamp() * 1000)


def bucket_text(bucket_ms: int, aggregation: str) -> str:
    local_dt = datetime.fromtimestamp(bucket_ms / 1000, ZoneInfo(APP_TIMEZONE))
    if aggregation == "daily":
        return local_dt.strftime("%Y-%m-%d")
    end_dt = local_dt.replace(minute=59, second=0, microsecond=0)
    return f"{local_dt:%Y-%m-%d %H:00}-{end_dt:%H:%M}"


def to_float(value: object) -> float | None:
    return float(value) if value is not None else None


def append_header(sheet: Worksheet, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def append_info_sheet(
    workbook: Workbook,
    payload: ExportAggregatesRequest,
    station_rows: list[MonitoringPost],
    start_local: datetime,
    end_local: datetime,
) -> None:
    sheet = workbook.active
    sheet.title = "info"
    rows = [
        ("aggregation", payload.aggregation),
        ("timezone", APP_TIMEZONE),
        ("start", start_local.strftime("%Y-%m-%d %H:%M:%S")),
        ("end", end_local.strftime("%Y-%m-%d %H:%M:%S")),
        ("stations", ", ".join(post.name or "Без названия" for post in station_rows)),
        ("device_types", ", ".join(payload.device_types)),
    ]
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 60


def base_query(model: type, station_ids: list[int], start_ms: int, end_ms: int):
    return (
        select(model, MonitoringPost.serial.label("post_serial"), MonitoringPost.name.label("post_name"))
        .join(MonitoringPost, MonitoringPost.id == model.monitoring_post_id)
        .where(
            model.monitoring_post_id.in_(station_ids),
            model.bucket_ms >= start_ms,
            model.bucket_ms <= end_ms,
        )
        .order_by(MonitoringPost.serial.asc(), model.bucket_ms.asc())
    )


def sort_gas_codes(codes: list[str]) -> list[str]:
    known_codes = [code for code in GAS_SUBSTANCE_ORDER if code in codes]
    other_codes = sorted(code for code in codes if code not in GAS_SUBSTANCE_ORDER)
    return known_codes + other_codes


def append_gas_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("gas")
    substance_rows = db.execute(
        select(model.substance_code)
        .where(
            model.monitoring_post_id.in_(station_ids),
            model.bucket_ms >= start_ms,
            model.bucket_ms <= end_ms,
            model.value_avg.is_not(None),
        )
        .distinct()
    )
    substance_codes = sort_gas_codes(list({*GAS_SUBSTANCE_ORDER, *(row.substance_code for row in substance_rows)}))
    append_header(sheet, ["station_serial", "station_name", "time", *substance_codes])

    rows = db.execute(
        base_query(model, station_ids, start_ms, end_ms)
        .where(model.value_avg.is_not(None))
        .order_by(
            MonitoringPost.serial.asc(),
            model.bucket_ms.asc(),
            model.substance_code.asc(),
        )
    )
    current_key = None
    current_values: dict[str, float | None] = {}
    current_serial = None
    current_bucket_ms = None

    for row in rows:
        item = row[0]
        key = (row.post_serial, item.bucket_ms)
        if current_key is not None and key != current_key:
            sheet.append([
                current_serial,
                current_name,
                bucket_text(current_bucket_ms, aggregation),
                *(current_values.get(code) for code in substance_codes),
            ])
            row_count += 1
            if row_count > MAX_EXPORT_ROWS:
                raise HTTPException(status_code=413, detail="Слишком большой экспорт. Уменьшите диапазон или выборку.")
            current_values = {}

        current_key = key
        current_serial = row.post_serial
        current_name = row.post_name
        current_bucket_ms = item.bucket_ms
        current_values[item.substance_code] = to_float(item.value_avg)

    if current_key is not None:
        sheet.append([
            current_serial,
            current_name,
            bucket_text(current_bucket_ms, aggregation),
            *(current_values.get(code) for code in substance_codes),
        ])
        row_count += 1
        if row_count > MAX_EXPORT_ROWS:
            raise HTTPException(status_code=413, detail="Слишком большой экспорт. Уменьшите диапазон или выборку.")

    return row_count


def append_dust_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("dust")
    append_header(sheet, [
        "station_serial",
        "station_name",
        "time",
        "pm1_mg_m3",
        "pm2_5_mg_m3",
        "pm10_mg_m3",
        "tsp_mg_m3",
    ])
    query = base_query(model, station_ids, start_ms, end_ms).where(
        or_(
            model.pm1_avg.is_not(None),
            model.pm2_avg.is_not(None),
            model.pm10_avg.is_not(None),
            model.tsp_avg.is_not(None),
        )
    )
    for row in db.execute(query):
        item = row[0]
        sheet.append([
            row.post_serial,
            row.post_name,
            bucket_text(item.bucket_ms, aggregation),
            to_float(item.pm1_avg),
            to_float(item.pm2_avg),
            to_float(item.pm10_avg),
            to_float(item.tsp_avg),
        ])
        row_count += 1
        if row_count > MAX_EXPORT_ROWS:
            raise HTTPException(status_code=413, detail="Слишком большой экспорт. Уменьшите диапазон или выборку.")
    return row_count


def append_meteo_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("meteo")
    append_header(sheet, [
        "station_serial",
        "station_name",
        "time",
        "air_temperature_c",
        "air_humidity_percent",
        "atmospheric_pressure",
        "wind_direction_deg",
        "wind_speed_m_s",
    ])
    query = base_query(model, station_ids, start_ms, end_ms).where(
        or_(
            model.air_temp_avg.is_not(None),
            model.air_hum_avg.is_not(None),
            model.atm_press_avg.is_not(None),
            model.hor_win_dir_avg.is_not(None),
            model.hor_win_spd_avg.is_not(None),
        )
    )
    for row in db.execute(query):
        item = row[0]
        sheet.append([
            row.post_serial,
            row.post_name,
            bucket_text(item.bucket_ms, aggregation),
            to_float(item.air_temp_avg),
            to_float(item.air_hum_avg),
            to_float(item.atm_press_avg),
            to_float(item.hor_win_dir_avg),
            to_float(item.hor_win_spd_avg),
        ])
        row_count += 1
        if row_count > MAX_EXPORT_ROWS:
            raise HTTPException(status_code=413, detail="Слишком большой экспорт. Уменьшите диапазон или выборку.")
    return row_count


def append_ivtm_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("ivtm")
    append_header(sheet, [
        "station_serial",
        "station_name",
        "time",
        "humidity_percent",
        "temperature_c",
    ])
    query = base_query(model, station_ids, start_ms, end_ms).where(
        or_(
            model.sensor_ivtm_hum_avg.is_not(None),
            model.sensor_ivtm_temp_avg.is_not(None),
        )
    )
    for row in db.execute(query):
        item = row[0]
        sheet.append([
            row.post_serial,
            row.post_name,
            bucket_text(item.bucket_ms, aggregation),
            to_float(item.sensor_ivtm_hum_avg),
            to_float(item.sensor_ivtm_temp_avg),
        ])
        row_count += 1
        if row_count > MAX_EXPORT_ROWS:
            raise HTTPException(status_code=413, detail="Слишком большой экспорт. Уменьшите диапазон или выборку.")
    return row_count


def append_profile_levels_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("profile_levels")
    append_header(sheet, ["station_serial", "station_name", "time", "height_m", "temperature_c"])
    rows = db.execute(
        base_query(model, station_ids, start_ms, end_ms)
        .where(model.temperature_avg.is_not(None))
        .order_by(
            MonitoringPost.serial.asc(),
            model.bucket_ms.asc(),
            model.height.asc(),
        )
    )
    for row in rows:
        item = row[0]
        sheet.append([
            row.post_serial,
            row.post_name,
            bucket_text(item.bucket_ms, aggregation),
            to_float(item.height),
            to_float(item.temperature_avg),
        ])
        row_count += 1
        if row_count > MAX_EXPORT_ROWS:
            raise HTTPException(status_code=413, detail="Слишком большой экспорт. Уменьшите диапазон или выборку.")
    return row_count


def append_profile_inversion_sheet(
    workbook: Workbook,
    db: Session,
    model: type,
    station_ids: list[int],
    start_ms: int,
    end_ms: int,
    aggregation: str,
    row_count: int,
) -> int:
    sheet = workbook.create_sheet("profile_inversion")
    append_header(sheet, [
        "station_serial",
        "station_name",
        "time",
        "inversion_power",
        "inversion_lower_m",
        "inversion_upper_m",
        "inversion_delta_t_c",
    ])
    query = base_query(model, station_ids, start_ms, end_ms).where(
        or_(
            model.inversion_power_avg.is_not(None),
            model.inversion_lower_avg.is_not(None),
            model.inversion_upper_avg.is_not(None),
            model.inversion_delta_t_avg.is_not(None),
        )
    )
    for row in db.execute(query):
        item = row[0]
        sheet.append([
            row.post_serial,
            row.post_name,
            bucket_text(item.bucket_ms, aggregation),
            to_float(item.inversion_power_avg),
            to_float(item.inversion_lower_avg),
            to_float(item.inversion_upper_avg),
            to_float(item.inversion_delta_t_avg),
        ])
        row_count += 1
        if row_count > MAX_EXPORT_ROWS:
            raise HTTPException(status_code=413, detail="Слишком большой экспорт. Уменьшите диапазон или выборку.")
    return row_count


def filename_datetime_part(value: datetime) -> str:
    return value.strftime("%Y-%m-%d_%H-%M")


def build_filename(
    payload: ExportAggregatesRequest,
    start_local: datetime,
    end_local: datetime,
) -> str:
    period_part = f"{filename_datetime_part(start_local)}_to_{filename_datetime_part(end_local)}"
    return f"eco_export_{payload.aggregation}_{period_part}.xlsx"


@router.post("/aggregates")
def export_aggregates(
    payload: ExportAggregatesRequest,
    db: Session = Depends(get_db),
    _current_user: User = Depends(get_current_user),
) -> Response:
    start_local = parse_local_datetime(payload.start, is_end=False)
    end_local = parse_local_datetime(payload.end, is_end=True)
    if end_local < start_local:
        raise HTTPException(status_code=422, detail="end must be greater than or equal to start")

    if not payload.device_types:
        raise HTTPException(status_code=422, detail="Select at least one device type")

    duplicate_device_types = len(set(payload.device_types)) != len(payload.device_types)
    if duplicate_device_types:
        raise HTTPException(status_code=422, detail="Device types must be unique")

    if payload.station_ids is None:
        station_rows = db.scalars(
            select(MonitoringPost)
            .where(MonitoringPost.is_confirmed.is_(True))
            .order_by(MonitoringPost.serial.asc())
        ).all()
    else:
        if not payload.station_ids:
            raise HTTPException(status_code=422, detail="Select at least one station")
        station_rows = db.scalars(
            select(MonitoringPost)
            .where(
                MonitoringPost.id.in_(payload.station_ids),
                MonitoringPost.is_confirmed.is_(True),
            )
            .order_by(MonitoringPost.serial.asc())
        ).all()
        found_ids = {row.id for row in station_rows}
        missing_ids = set(payload.station_ids) - found_ids
        if missing_ids:
            raise HTTPException(status_code=404, detail="One or more stations were not found")

    if not station_rows:
        raise HTTPException(status_code=422, detail="No confirmed stations available for export")

    start_ms = to_epoch_ms(start_local)
    end_ms = to_epoch_ms(end_local)
    station_ids = [row.id for row in station_rows]

    workbook = Workbook()
    append_info_sheet(workbook, payload, station_rows, start_local, end_local)
    row_count = 0

    if payload.aggregation == "hourly":
        models = {
            "gas": CaggGasHourly,
            "dust": CaggDustHourly,
            "meteo": CaggMeteoHourly,
            "ivtm": CaggIvtmHourly,
            "profile_levels": CaggProfileLevelsHourly,
            "profile_inversion": CaggProfileInversionHourly,
        }
    else:
        models = {
            "gas": CaggGasDaily,
            "dust": CaggDustDaily,
            "meteo": CaggMeteoDaily,
            "ivtm": CaggIvtmDaily,
            "profile_levels": CaggProfileLevelsDaily,
            "profile_inversion": CaggProfileInversionDaily,
        }

    if "gas" in payload.device_types:
        row_count = append_gas_sheet(workbook, db, models["gas"], station_ids, start_ms, end_ms, payload.aggregation, row_count)
    if "dust" in payload.device_types:
        row_count = append_dust_sheet(workbook, db, models["dust"], station_ids, start_ms, end_ms, payload.aggregation, row_count)
    if "meteo" in payload.device_types:
        row_count = append_meteo_sheet(workbook, db, models["meteo"], station_ids, start_ms, end_ms, payload.aggregation, row_count)
    if "ivtm" in payload.device_types:
        row_count = append_ivtm_sheet(workbook, db, models["ivtm"], station_ids, start_ms, end_ms, payload.aggregation, row_count)
    if "profile" in payload.device_types:
        row_count = append_profile_levels_sheet(
            workbook,
            db,
            models["profile_levels"],
            station_ids,
            start_ms,
            end_ms,
            payload.aggregation,
            row_count,
        )
        row_count = append_profile_inversion_sheet(
            workbook,
            db,
            models["profile_inversion"],
            station_ids,
            start_ms,
            end_ms,
            payload.aggregation,
            row_count,
        )

    output = BytesIO()
    workbook.save(output)
    filename = build_filename(payload, start_local, end_local)
    headers = {
        "Content-Disposition": f"attachment; filename={filename}; filename*=UTF-8''{filename}",
    }
    return Response(content=output.getvalue(), media_type=XLSX_MEDIA_TYPE, headers=headers)
